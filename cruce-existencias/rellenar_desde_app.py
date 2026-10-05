"""Rellena y verifica la Ficha de Descarte PP 2026 usando los datos de FOLIO que ya guardó
la app de Descarte en Ficha_Descarte_PowerApps.xlsx (tblFichaDescarte).

Uso: python rellenar_desde_app.py Ficha.xlsx app.pkl Salida.xlsx
Verde = dato copiado de la app (FOLIO). Naranjo = deducido por la serie del código (mismo título).
"""
import re
import sys
from datetime import datetime

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill

VERDE = PatternFill("solid", fgColor="E2EFDA")
NARANJO = PatternFill("solid", fgColor="FCE4D6")
EX = ["Pregrado A - Existencia", "Pregrado F - Existencia", "Posgrado - Existencia", "Viña - Existencia"]
EX_APP = ["Pregrado  A - Existencia", "Pregrado  F - Existencia", "Posgrado - Existencia", "Viña - Existencia"]


def txt(v):
    if v is None or (isinstance(v, float) and v != v):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, datetime):
        return v.strftime("%d-%m-%Y")
    return str(v).strip()


def norm_cod(v):
    s = txt(v)
    return s.zfill(6) if s.isdigit() and len(s) < 6 else s


def anio_de(copia):
    m = re.search(r"(1[89]\d\d|20\d\d)", copia)
    return int(m.group(1)) if m else ""


def biblio_de(ubic):
    b = ubic.split(" - ")[0].strip()
    return b


def col_propia(ubic):
    for clave, c in (("Edif. A", EX[0]), ("Edif. F", EX[1]), ("Postgrado", EX[2]), ("Viña", EX[3])):
        if clave.lower() in ubic.lower():
            return c
    return None


def datos_app(a):
    ubic = a["Biblioteca-Ubicacion-Colección"]
    propia = col_propia(ubic)
    d = {
        "Item ingresado en la base de biblioteca": "SI",
        "HRID": int(a["HRID"].lstrip("0") or 0) if a["HRID"].isdigit() else a["HRID"],
        "Copia": a["copia"],
        "Tipo de Material": a["Tipo de Material"],
        "Biblioteca-Ubicacion-Colección": ubic.replace(" - ", " "),
        "Título": a["Título"],
        "Autor": a["Autor"],
        "Idioma": a["Idioma"],
        "Vinculado UAI (SI/NO)": a["Vinculado UAI (SI/NO)"] or "NO",
        "Fecha registro (ingresado en la base)": datetime.strptime(a["Fecha registro (ingresado en la base)"], "%d-%m-%Y")
        if re.match(r"\d\d-\d\d-\d{4}$", a["Fecha registro (ingresado en la base)"]) else "",
        "Unidad Academica o Centro de Costo *": a["Unidad Academica o Centro de Costo *"],
        "Año de edicion": int(a["Año de edición"]) if a["Año de edición"].isdigit() else "",
        "Unidad de descarte": biblio_de(ubic),
    }
    # Existencias: 1 en la biblioteca propia; en las otras, 0 si el título no tiene ejemplares ahí.
    # Si el título sí tiene ejemplares en otra biblioteca, no se sabe cuántos son del mismo número → se deja vacío.
    if d["Año de edicion"] == "":
        d["Año de edicion (deducido)"] = anio_de(a["copia"])
    otras_desconocidas = []
    for c, ca in zip(EX, EX_APP):
        if c == propia:
            d[c] = 1
        elif txt(a[ca]) in ("", "0"):
            d[c] = 0
        else:
            otras_desconocidas.append(c)
    return d, otras_desconocidas


def igual(k, ficha, nuevo):
    a, b = txt(ficha), txt(nuevo)
    if a.lower() == b.lower() or {a, b} <= {"", "0"}:
        return True
    if k == "Unidad de descarte":
        return b.lower() in a.lower() or a.lower() in b.lower()
    try:
        return float(a) == float(b)
    except ValueError:
        return re.sub(r"\W", "", a.lower()) == re.sub(r"\W", "", b.lower())


def main(ficha, pkl, salida):
    app = pd.read_pickle(pkl).fillna("")
    app = app.apply(lambda s: s.str.strip())
    por_cod = {norm_cod(r["Codigo de Barra"]): r for r in app.to_dict("records")}
    wb = openpyxl.load_workbook(ficha)
    ws = wb["Ficha descarte - Base"]
    col = {txt(c.value): c.column for c in ws[1] if c.value}

    # Serie 36729-…: todas las que la app trae son del mismo título → plantilla para las que no trae
    serie = {}
    for c, a in por_cod.items():
        if "-" in c:
            serie.setdefault(c.split("-")[0], a)

    dif, notas, cont = [], [], {"app": 0, "serie": 0, "sin_dato": 0, "verificadas": 0}
    for r in range(3, ws.max_row + 1):
        cod = norm_cod(ws.cell(r, col["Codigo de Barra"]).value)
        if not cod:
            continue
        vacia = txt(ws.cell(r, col["HRID"]).value) == ""
        a = por_cod.get(cod)
        relleno = VERDE
        if a is None and vacia and "-" in cod and cod.split("-")[0] in serie:
            a = dict(serie[cod.split("-")[0]])
            a.update({"copia": "", "Fecha registro (ingresado en la base)": "", "Año de edición": ""})
            relleno = NARANJO
            notas.append([r, cod, "Deducido por la serie del código (mismo título). Falta Copia/volumen, fecha y año."])
            cont["serie"] += 1
        elif a is None:
            if not vacia and all(txt(ws.cell(r, col[c]).value) == "" for c in EX):
                notas.append([r, cod, "Existencias vacías y este código no está en los datos de la app: completar en la sesión local."])
            if vacia:
                cont["sin_dato"] += 1
                notas.append([r, cod, "Sin datos de FOLIO disponibles: completar en la sesión local."])
            continue
        elif vacia:
            cont["app"] += 1
        else:
            cont["verificadas"] += 1
        d, desconocidas = datos_app(a)
        for k, v in d.items():
            deducido = k.endswith(" (deducido)")
            k = k.replace(" (deducido)", "")
            celda = ws.cell(r, col[k])
            if txt(celda.value) == "":
                if v not in ("", None):
                    celda.value = v
                    celda.fill = relleno
                    if isinstance(v, datetime):
                        celda.number_format = "dd-mm-yyyy"
            elif not vacia and not deducido and txt(v) != "" and not igual(k, celda.value, v) and k not in (
                    "Item ingresado en la base de biblioteca", "Vinculado UAI (SI/NO)",
                    "Unidad Academica o Centro de Costo *", "Fecha registro (ingresado en la base)"):
                dif.append([r, cod, k, txt(celda.value), txt(v)])
        for c in desconocidas:
            notas.append([r, cod, f"{c}: el título tiene ejemplares en esa biblioteca; confirmar cuántos son de este número."])

    wd = wb.create_sheet("Diferencias")
    wd.append(["Fila", "Codigo de Barra", "Columna", "Dice la ficha", "Dice FOLIO (app Descarte)"])
    for x in dif:
        wd.append(x)
    wn = wb.create_sheet("Pendientes")
    wn.append(["Fila", "Codigo de Barra", "Qué falta"])
    for x in notas:
        wn.append(x)
    for h, anchos in ((wd, (8, 14, 34, 45, 45)), (wn, (8, 14, 95))):
        for c in h[1]:
            c.font = Font(bold=True)
        for i, w in enumerate(anchos):
            h.column_dimensions["ABCDE"[i]].width = w
        h.freeze_panes = "A2"
        h.auto_filter.ref = h.dimensions
    wb.save(salida)
    print(cont, "diferencias:", len(dif), "pendientes:", len(notas))


if __name__ == "__main__":
    main(*sys.argv[1:4])
