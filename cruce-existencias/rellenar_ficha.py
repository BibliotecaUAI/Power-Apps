"""Rellena y verifica la Ficha de Descarte PP 2026 con los datos de FOLIO.

Uso:  python rellenar_ficha.py  Ficha.xlsx  Datos_FOLIO.xlsx  Salida.xlsx
      (Datos_FOLIO = la tabla "Cruce_FOLIO" copiada desde Power BI Desktop y pegada en Excel;
       también acepta el CSV con separador "|".)

- Filas sin datos (solo código de barra): rellena todas las columnas que vienen de FOLIO.
- Filas con datos: rellena las celdas de FOLIO que estén vacías y anota en la hoja
  "Diferencias" lo que no coincide con FOLIO (no cambia valores ya escritos).
- Existencias (T a W): 1 en la biblioteca del propio ejemplar; en las otras, la cantidad
  de copias de ese mismo número (HRID + volumen), sin tope.
- Columnas manuales: no se tocan.
Las celdas rellenadas quedan en verde.
"""
import ast
import re
import sys
from datetime import datetime

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill

HOJA = "Ficha descarte - Base"
FILA_INICIO = 3
VERDE = PatternFill("solid", fgColor="E2EFDA")
BIBLIOTECAS = [  # (texto en FOLIO, encabezado de la ficha)
    ("Edif. A", "Pregrado A - Existencia"),
    ("Edif. F", "Pregrado F - Existencia"),
    ("Postgrado", "Posgrado - Existencia"),
    ("Viña", "Viña - Existencia"),
]
COLS_EX = [c for _, c in BIBLIOTECAS]


def txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def norm_volumen(v):
    s = re.sub(r"\s+", "", txt(v).upper())
    return re.sub(r"(C|COP|COPIA|EJ|EJEMPLAR)\.?\d+$", "", s)


def biblioteca_col(texto):
    for clave, col in BIBLIOTECAS:
        if clave.lower() in texto.lower():
            return col
    return None


def autor(contrib):
    try:
        lst = ast.literal_eval(contrib) if contrib else []
        return "; ".join(c["name"] for c in lst if c.get("name")) or "0"
    except (ValueError, SyntaxError, TypeError, AttributeError):
        return contrib or "0"


def idioma(langs):
    try:
        return ", ".join(ast.literal_eval(langs)) if langs else ""
    except (ValueError, SyntaxError):
        return langs.strip("[]").replace("'", "")


def igual(a, b):
    a, b = txt(a), txt(b)
    if a.lower() == b.lower():
        return True
    try:
        return float(a) == float(b)
    except ValueError:
        return re.sub(r"\W", "", a.lower()) == re.sub(r"\W", "", b.lower())


def main(ficha, csv, salida):
    if str(csv).lower().endswith((".xlsx", ".xls")):
        ej = pd.read_excel(csv, dtype=str, keep_default_na=False)
    else:
        ej = pd.read_csv(csv, sep="|", dtype=str, encoding="utf-8-sig", keep_default_na=False)
    ej.columns = [c.strip("[] ").split("[")[-1] for c in ej.columns]
    ej = ej.apply(lambda s: s.str.strip())
    ej["fecha"] = ej["fecha"].str[:10]
    ej["vol"] = ej["copia"].map(norm_volumen)
    ej["col"] = ej["biblioteca"].map(biblioteca_col)
    conteo = (ej.dropna(subset=["col"]).groupby(["hrid", "vol", "col"])["codigo"].nunique()
              .unstack("col").reindex(columns=COLS_EX).fillna(0).astype(int))
    por_cod = {c: r for c, r in zip(ej["codigo"], ej.to_dict("records"))}

    wb = openpyxl.load_workbook(ficha)
    ws = wb[HOJA]
    col = {txt(c.value): c.column for c in ws[1] if c.value}

    diferencias, resumen = [], {"rellenadas": 0, "verificadas": 0, "no_encontradas": 0}
    for r in range(FILA_INICIO, ws.max_row + 1):
        cod = txt(ws.cell(r, col["Codigo de Barra"]).value)
        if not cod:
            continue
        it = por_cod.get(cod) or (por_cod.get(cod.zfill(6)) if cod.isdigit() else None)
        if it is None:
            resumen["no_encontradas"] += 1
            diferencias.append([r, cod, "Codigo de Barra", cod, "", "No se encontró en FOLIO"])
            continue
        ex = conteo.loc[(it["hrid"], it["vol"])] if (it["hrid"], it["vol"]) in conteo.index else None
        propia = biblioteca_col(it["biblioteca"])
        datos = {
            "Item ingresado en la base de biblioteca": "SI",
            "HRID": int(it["hrid"]) if it["hrid"].isdigit() else it["hrid"],
            "Copia": it["copia"],
            "Tipo de Material": it["tipo"],
            "Biblioteca-Ubicacion-Colección": f'{it["biblioteca"]} {it["ubicacion"]}'.strip(),
            "Título": it["titulo"],
            "Autor": autor(it["autor"]),
            "Idioma": idioma(it["idioma"]),
            "Vinculado UAI (SI/NO)": "NO",
            "Fecha registro (ingresado en la base)": datetime.strptime(it["fecha"], "%Y-%m-%d") if it["fecha"] else "",
            "Unidad Academica o Centro de Costo *": it["unidad"] or "None",
            "Año de edicion": int(it["anio"]) if it["anio"].isdigit() else it["anio"],
            "Unidad de descarte": it["biblioteca"],
        }
        for c in COLS_EX:
            datos[c] = 1 if c == propia else (int(ex[c]) if ex is not None else 0)
        vacia = all(txt(ws.cell(r, col[k]).value) == "" for k in datos if k in col)
        if vacia and txt(ws.cell(r, col["Cruce Archivo activo Fijo (Finanzas)"]).value) == "":
            datos["Cruce Archivo activo Fijo (Finanzas)"] = "No"
        resumen["rellenadas" if vacia else "verificadas"] += 1
        for k, v in datos.items():
            celda = ws.cell(r, col[k])
            if txt(celda.value) == "":
                if v not in ("", None):
                    celda.value = v
                    celda.fill = VERDE
                    if k == "Fecha registro (ingresado en la base)":
                        celda.number_format = "dd-mm-yyyy"
            elif not igual(celda.value, v) and k not in ("Vinculado UAI (SI/NO)",
                                                          "Unidad Academica o Centro de Costo *"):
                diferencias.append([r, cod, k, txt(celda.value), txt(v), "Distinto a FOLIO"])

    wd = wb.create_sheet("Diferencias")
    wd.append(["Fila", "Codigo de Barra", "Columna", "Dice la ficha", "Dice FOLIO", "Observación"])
    for d in diferencias:
        wd.append(d)
    for c in wd[1]:
        c.font = Font(bold=True)
    for letra, ancho in zip("ABCDEF", (8, 14, 34, 40, 40, 26)):
        wd.column_dimensions[letra].width = ancho
    wd.freeze_panes = "A2"
    wd.auto_filter.ref = wd.dimensions
    wb.save(salida)
    print(f"Filas rellenadas: {resumen['rellenadas']} · verificadas: {resumen['verificadas']} · "
          f"no encontradas en FOLIO: {resumen['no_encontradas']} · diferencias: {len(diferencias)}")


if __name__ == "__main__":
    main(*sys.argv[1:4])
